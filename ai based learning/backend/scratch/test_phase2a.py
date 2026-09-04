import sys
from pathlib import Path
import json

backend_path = Path(r"c:\Users\vishr\Downloads\sih\backend")
sys.path.insert(0, str(backend_path))

from fastapi.testclient import TestClient
from main import app
import fitz  # PyMuPDF
from pptx import Presentation
import docx

client = TestClient(app)

def create_sample_pdf(file_path: Path):
    """Create a multi-page PDF using PyMuPDF."""
    doc = fitz.open()
    
    # Page 1
    page1 = doc.new_page()
    page1.insert_text((50, 50), "National Accounts Statistics Manual 2026\nChapter 1: Gross Domestic Product (GDP) Compilation Methods.")
    
    # Page 2
    page2 = doc.new_page()
    page2.insert_text((50, 50), "Chapter 2: Gross Value Added (GVA) by Economic Activity.\nBasic Prices vs Producer Prices formula.")
    
    doc.save(str(file_path))
    doc.close()

def create_sample_pptx(file_path: Path):
    """Create a multi-slide PPTX using python-pptx."""
    prs = Presentation()
    
    # Slide 1
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Labour Force Survey (PLFS) Overview"
    slide1.placeholders[1].text = "Quarterly Bulletin & Annual Report Indicators."
    
    # Slide 2
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Worker Population Ratio (WPR) & LFPR"
    slide2.placeholders[1].text = "1. Usual Principal Status (UPS)\n2. Current Weekly Status (CWS)"
    
    prs.save(str(file_path))

def create_sample_docx(file_path: Path):
    """Create a structured DOCX using python-docx."""
    doc = docx.Document()
    doc.add_heading("Index of Industrial Production (IIP)", level=1)
    doc.add_paragraph("The IIP is a composite indicator that measures short-term changes in the volume of production of a basket of industrial products.")
    
    doc.add_heading("Weighting Diagram", level=2)
    doc.add_paragraph("Base year 2011-12 weighting diagram covers Mining (14.37%), Manufacturing (77.63%), and Electricity (7.99%).")
    
    doc.save(str(file_path))

def run_tests():
    print("=== RUNNING PHASE 2A TEXT EXTRACTION VERIFICATION ===")
    tmp_dir = backend_path / "data" / "tmp_test"
    tmp_dir.mkdir(parents=True, exist_ok=True)

    # 1. Verify PDF Text Extraction
    pdf_path = tmp_dir / "test_gdp_manual.pdf"
    create_sample_pdf(pdf_path)
    print("\n1. Testing PDF text extraction via POST /api/documents/upload ...")
    with open(pdf_path, "rb") as f:
        res = client.post("/api/documents/upload", files={"file": ("test_gdp_manual.pdf", f, "application/pdf")})
    
    print(f"Status Code: {res.status_code}")
    print(f"Upload Response: {res.json()}")
    assert res.status_code == 201
    doc_id_pdf = res.json()["data"]["document_id"]
    assert res.json()["data"]["status"] in ["TEXT_EXTRACTED", "CHUNKED"]
    assert res.json()["data"]["text_blocks"] >= 2

    # Verify Preview API for PDF
    print("\n1b. Testing GET /api/documents/{document_id}/preview for PDF ...")
    res_prev = client.get(f"/api/documents/{doc_id_pdf}/preview")
    print(f"Status Code: {res_prev.status_code}")
    preview_data = res_prev.json()
    print(f"Extracted blocks count: {len(preview_data['content'])}")
    assert len(preview_data['content']) >= 2
    assert "Gross Domestic Product" in preview_data['content'][0]['text']
    assert preview_data['content'][0]['location'] == "Page 1"
    assert preview_data['content'][1]['location'] == "Page 2"

    # 2. Verify PPTX Text Extraction
    pptx_path = tmp_dir / "test_plfs_presentation.pptx"
    create_sample_pptx(pptx_path)
    print("\n2. Testing PPTX text extraction via POST /api/documents/upload ...")
    with open(pptx_path, "rb") as f:
        res = client.post("/api/documents/upload", files={"file": ("test_plfs_presentation.pptx", f, "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
    
    print(f"Status Code: {res.status_code}")
    print(f"Upload Response: {res.json()}")
    assert res.status_code == 201
    doc_id_pptx = res.json()["data"]["document_id"]
    assert res.json()["data"]["status"] in ["TEXT_EXTRACTED", "CHUNKED"]

    res_prev = client.get(f"/api/documents/{doc_id_pptx}/preview")
    preview_data = res_prev.json()
    assert len(preview_data['content']) >= 2
    assert preview_data['content'][0]['location'] == "Slide 1"
    assert "Labour Force Survey" in preview_data['content'][0]['text']

    # 3. Verify DOCX Text Extraction
    docx_path = tmp_dir / "test_iip_handbook.docx"
    create_sample_docx(docx_path)
    print("\n3. Testing DOCX text extraction via POST /api/documents/upload ...")
    with open(docx_path, "rb") as f:
        res = client.post("/api/documents/upload", files={"file": ("test_iip_handbook.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    
    print(f"Status Code: {res.status_code}")
    print(f"Upload Response: {res.json()}")
    assert res.status_code == 201
    doc_id_docx = res.json()["data"]["document_id"]
    assert res.json()["data"]["status"] in ["TEXT_EXTRACTED", "CHUNKED"]

    res_prev = client.get(f"/api/documents/{doc_id_docx}/preview")
    preview_data = res_prev.json()
    assert len(preview_data['content']) >= 2
    assert "Index of Industrial Production" in preview_data['content'][0]['text']

    # 4. Verify Phase 1 Health & AI Diagnostic Endpoints Still Function Flawlessly
    print("\n4. Verifying Phase 1 Health and AI endpoints ...")
    h_res = client.get("/api/health")
    assert h_res.status_code == 200
    ai_res = client.post("/api/ai/test", json={"prompt": "Verify Phase 2A backward compatibility"})
    assert ai_res.status_code == 200

    print("\nPHASE 2A DOCUMENT EXTRACTION TESTS PASSED SUCCESSFULLY! [SUCCESS]")

if __name__ == "__main__":
    run_tests()
